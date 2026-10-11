"""Read-only cohort audit; does not refit or replace scientific outputs."""
from datetime import datetime, timezone
import gzip
import hashlib
import io
import json
from pathlib import Path
import xml.etree.ElementTree as ET

import httpx
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'results/validation/github/cohort'


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    geo = json.loads((ROOT/'results/phase4/geo_review/GSE284723.json').read_text())
    workbook = ROOT/'data/raw/GSE284723/GSE284723_rawCounts.xlsx'
    counts = pd.read_excel(workbook, sheet_name='raw counts')
    metadata = pd.read_csv(ROOT/'data/processed/GSE284723/sample_metadata.csv')
    assert list(counts.columns[1:]) == metadata.count_column.tolist()
    assert set(metadata.accession) == {s['accession'] for s in geo['samples']}
    assert metadata.condition.value_counts().to_dict() == {'Control':5, 'mKO':4}
    matrix = ROOT/'data/raw/GSE284723/GSE284723_series_matrix.txt.gz'
    lines = gzip.decompress(matrix.read_bytes()).decode().splitlines()
    relations = [line for line in lines if line.startswith('!Sample_relation')]
    record = {'checked_utc':datetime.now(timezone.utc).isoformat(), 'status':'unresolved',
              'design_mice':10, 'deposited_libraries':9, 'paper_figure4_mice':8,
              'mapping':metadata.to_dict('records'), 'raw_columns':counts.columns.tolist(),
              'raw_rows':len(counts), 'summary_rows':int(counts.EnsemblID.str.startswith('__').sum()),
              'deposited_relations':relations, 'sources':[],
              'contrast':'mKO versus Control; ~ condition; all nine deposited libraries retained',
              'decision':'Exclude third-study findings from validated headline evidence; retain exploratory outputs unchanged.',
              'inference':'Absent Crep4 label is not evidence of an exclusion or its reason. Unique BioSamples do not establish animal independence.',
              'source_hashes':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in (matrix,workbook)},
              'reviewed_scripts':['scripts/discovery/prepare_third.py','scripts/discovery/third_deseq2.R'],
              'previous_logs':['results/validation/phase5/reproduction_04.log'],
              'article':'https://pmc.ncbi.nlm.nih.gov/articles/PMC12673120/'}
    with httpx.Client(timeout=60, follow_redirects=True) as client:
        url = 'https://api.figshare.com/v2/articles/30090307'
        response = client.get(url)
        response.raise_for_status()
        item = response.json()
        (OUT/'figshare_metadata.json').write_text(json.dumps(item,indent=2)+'\n')
        file = item['files'][0]
        data = client.get(file['download_url'])
        data.raise_for_status()
        book = openpyxl.load_workbook(io.BytesIO(data.content), read_only=True, data_only=True)
        matches = {}
        for sheet in book:
            hits = []
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value,str) and any(term in cell.value.lower() for term in ('rna','crep','cren','exclude','outlier','seq')):
                        hits.append({'cell':cell.coordinate,'value':cell.value[:300]})
            matches[sheet.title] = {'rows':sheet.max_row,'columns':sheet.max_column,'matches':hits}
        record['figshare'] = {'url':url, 'license':item['license'], 'download':file['download_url'],
                             'sha256':hashlib.sha256(data.content).hexdigest(),'sheets':matches}
        record['sources'].append({'url':url,'status':response.status_code})
        url = 'https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi?db=biosample&id='+','.join('SAMN'+str(n) for n in range(45897314,45897323))
        try:
            response = client.get(url)
            response.raise_for_status()
            tree = ET.fromstring(response.content)
            samples = tree.findall('.//BioSample')
            if len(samples) != 9:
                raise ValueError(f'Expected nine BioSamples; returned {len(samples)}')
            (OUT/'biosamples.xml').write_bytes(response.content)
            record['biosamples'] = [{'accession':s.get('accession'), 'attributes':[{ 'name':a.get('attribute_name'),'value':a.text} for a in s.findall('.//Attribute')]} for s in samples]
            record['sources'].append({'url':url,'status':response.status_code})
        except Exception as error:
            record['sources'].append({'url':url,'status':'unavailable','error':str(error)})
    (OUT/'audit.json').write_text(json.dumps(record,indent=2)+'\n')
    print(json.dumps(record,indent=2))


if __name__ == '__main__':
    main()
